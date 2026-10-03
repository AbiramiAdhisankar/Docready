"""Applies FormattingSettings to a DOCX using python-docx plus small, careful XML edits.

The input bytes are only read; a brand-new DOCX is returned. Nothing outside the
selected options is changed, and content (text, images, tables, links) is preserved
unless the user explicitly turned the matching "preserve" switch off.
"""
import io
from typing import Optional

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt
from docx.text.paragraph import Paragraph

from models.formatting_settings import MARGIN_PRESETS, PAGE_SIZES_INCHES, FormattingSettings
from services.docx_helpers import (
    has_drawing, heading_level, is_blank, is_centered, is_chapter_label, is_list_paragraph,
    is_short_line, is_special_style, local, runs_of,
)
from services.document_analyzer import spacing_label

ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
         "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}
ALIGN_LABEL = {"left": "Left", "center": "Center", "right": "Right", "justify": "Justified"}
THEME_ATTRS = ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme")


# ---- low-level run/style helpers -----------------------------------------------------
def _set_fonts(rpr, name: str) -> None:
    rfonts = rpr.get_or_add_rFonts()
    for attr in THEME_ATTRS:  # theme attributes would override the explicit names
        if qn(attr) in rfonts.attrib:
            del rfonts.attrib[qn(attr)]
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)


def _set_size(rpr, points: float) -> None:
    half = str(int(round(points * 2)))
    for tag in ("w:sz", "w:szCs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            rpr.append(el)
        el.set(qn("w:val"), half)


def _font_run(run, family: Optional[str], size: Optional[float]) -> None:
    rpr = run._r.get_or_add_rPr()
    if family:
        _set_fonts(rpr, family)
    if size:
        _set_size(rpr, size)


def _font_style(style, family, size, bold=None) -> None:
    rpr = style.element.get_or_add_rPr()
    if family:
        _set_fonts(rpr, family)
    if size:
        _set_size(rpr, size)
    if bold is not None:
        style.font.bold = bold


class DocumentFormatter:
    def __init__(self, data: bytes, settings: FormattingSettings):
        self.doc = Document(io.BytesIO(data))
        self.s = settings
        self.summary: list[str] = []
        self.stats = {"blank_removed": 0, "images_removed": 0, "tables_removed": 0,
                      "links_removed": 0, "images_scaled": 0, "headings_formatted": 0,
                      "labels_aligned": 0, "centered_kept": 0}

    # ---- public ------------------------------------------------------------------------
    def run(self) -> tuple[bytes, dict]:
        self._apply_preserve_options()
        if self.s.remove_blank_paragraphs:
            self._remove_blank_paragraphs()
        self._format_styles()
        self._format_paragraphs()
        self._format_tables()
        self._apply_page_setup()
        self._scale_images()
        self._apply_headers_footers()
        self._summarise()
        out = io.BytesIO()
        self.doc.save(out)
        return out.getvalue(), {"summary": self.summary, "after": self._after(), "stats": self.stats}

    # ---- preserve switches ---------------------------------------------------------------
    def _apply_preserve_options(self) -> None:
        body = self.doc.element.body
        if not self.s.preserve_hyperlinks:
            for h in list(body.iter(qn("w:hyperlink"))):
                parent, idx = h.getparent(), h.getparent().index(h)
                for child in list(h):
                    for rs in child.findall(qn("w:rPr") + "/" + qn("w:rStyle")):
                        if rs.get(qn("w:val")) == "Hyperlink":
                            rs.getparent().remove(rs)
                    parent.insert(idx, child)
                    idx += 1
                parent.remove(h)
                self.stats["links_removed"] += 1
        if not self.s.preserve_images:
            for el in list(body.iter()):
                name = local(el)
                is_img = (name == "drawing" and el.xpath(".//pic:pic")) or \
                         (name == "pict" and any(local(x) == "imagedata" for x in el.iter()))
                if not is_img or el.getparent() is None:
                    continue
                run = next((a for a in el.iterancestors() if local(a) == "r"), None)
                if run is not None and not any((t.text or "").strip() for t in run.iter(qn("w:t"))):
                    run.getparent().remove(run)
                else:
                    el.getparent().remove(el)
                self.stats["images_removed"] += 1
        if not self.s.preserve_tables:
            for tbl in list(body.iterchildren(qn("w:tbl"))):
                body.remove(tbl)
                self.stats["tables_removed"] += 1

    def _remove_blank_paragraphs(self) -> None:
        body = self.doc.element.body
        for p in list(body.iterchildren(qn("w:p"))):
            if not is_blank(p):
                continue
            prev, nxt = p.getprevious(), p.getnext()
            prev_tbl = prev is not None and local(prev) == "tbl"
            nxt_tbl = nxt is not None and local(nxt) == "tbl"
            at_end = nxt is None or local(nxt) == "sectPr"
            if prev_tbl and (nxt_tbl or at_end):  # keeps tables apart / valid body ending
                continue
            body.remove(p)
            self.stats["blank_removed"] += 1

    # ---- styles ----------------------------------------------------------------------------
    def _format_styles(self) -> None:
        s, styles = self.s, self.doc.styles
        try:
            normal = styles["Normal"]
        except KeyError:
            normal = None
        if normal is not None:
            _font_style(normal, s.font_family if s.normalize_fonts else None,
                        s.font_size if s.normalize_font_size else None)
            pf = normal.paragraph_format
            if s.normalize_alignment:
                pf.alignment = ALIGN[s.alignment]
            if s.normalize_line_spacing:
                pf.line_spacing = s.line_spacing
            if s.normalize_paragraph_spacing:
                pf.space_before, pf.space_after = Pt(s.space_before), Pt(s.space_after)
        try:
            title = styles["Title"]
            _font_style(title, s.title.font_family, s.title.size, s.title.bold)
            title.paragraph_format.alignment = ALIGN[s.title.alignment]
        except KeyError:
            pass
        for level, cfg in ((1, s.heading1), (2, s.heading2), (3, s.heading3)):
            try:
                st = styles[f"Heading {level}"]
            except KeyError:
                continue
            _font_style(st, self._heading_family(cfg), cfg.size, cfg.bold)
            if s.normalize_alignment and st.paragraph_format.alignment is None:
                st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT  # don't inherit "justify"

    def _heading_family(self, cfg) -> Optional[str]:
        return cfg.font_family or (self.s.font_family if self.s.normalize_fonts else None)

    # ---- paragraphs --------------------------------------------------------------------------
    def _format_paragraphs(self) -> None:
        s = self.s
        for p in self.doc.paragraphs:
            lvl = heading_level(p)
            if lvl == 0:
                for r in runs_of(p):
                    _font_run(r, s.title.font_family, s.title.size)
                    r.font.bold = s.title.bold
                p.alignment = ALIGN[s.title.alignment]
                self.stats["headings_formatted"] += 1
                continue
            if lvl:
                cfg = {1: s.heading1, 2: s.heading2, 3: s.heading3}.get(lvl)
                for r in runs_of(p):
                    if cfg:
                        _font_run(r, self._heading_family(cfg), cfg.size)
                        r.font.bold = cfg.bold
                    else:
                        _font_run(r, s.font_family if s.normalize_fonts else None, None)
                self.stats["headings_formatted"] += 1 if cfg else 0
                continue

            for r in runs_of(p):
                _font_run(r, s.font_family if s.normalize_fonts else None,
                          s.font_size if s.normalize_font_size else None)
                if s.font_style == "bold":
                    r.font.bold = True
                elif s.font_style == "italic":
                    r.font.italic = True
                elif s.font_style == "underline":
                    r.font.underline = True

            image_only = has_drawing(p._p) and not p.text.strip()
            if image_only or is_special_style(p):
                continue
            text = p.text
            is_label = s.chapter_labels.enabled and is_chapter_label(text)
            keep_center = (not is_label and s.keep_centered_lines and is_short_line(text)
                           and is_centered(p))
            if s.normalize_alignment:
                if is_label:
                    p.alignment = ALIGN[s.chapter_labels.alignment]
                    self.stats["labels_aligned"] += 1
                elif keep_center:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER  # explicit, so the Normal style can't override it
                    self.stats["centered_kept"] += 1
                else:
                    p.alignment = ALIGN[s.alignment]
            if s.normalize_line_spacing:
                p.paragraph_format.line_spacing = s.line_spacing
            if s.normalize_paragraph_spacing:
                p.paragraph_format.space_before = Pt(s.space_before)
                p.paragraph_format.space_after = Pt(s.space_after)
                left = p.paragraph_format.left_indent
                if is_label or keep_center:
                    p.paragraph_format.first_line_indent = Inches(0)
                elif not is_list_paragraph(p._p) and not (left and left > 0):
                    p.paragraph_format.first_line_indent = Inches(s.first_line_indent)

    def _format_tables(self) -> None:
        """Tables keep their own layout; only font family/size are normalised."""
        s = self.s
        if not (s.normalize_fonts or s.normalize_font_size):
            return
        for tbl in self.doc.element.body.iterchildren(qn("w:tbl")):
            for p_el in tbl.iter(qn("w:p")):
                for r in runs_of(Paragraph(p_el, self.doc)):
                    _font_run(r, s.font_family if s.normalize_fonts else None,
                              s.font_size if s.normalize_font_size else None)

    # ---- page setup ---------------------------------------------------------------------------
    def _margins(self):
        s = self.s
        if s.margin_preset == "custom":
            m = s.custom_margins
            return m.top, m.bottom, m.left, m.right
        return MARGIN_PRESETS[s.margin_preset]

    def _apply_page_setup(self) -> None:
        s = self.s
        w, h = PAGE_SIZES_INCHES[s.page_size]
        if s.orientation == "landscape":
            w, h = h, w
        top, bottom, left, right = self._margins()
        for sec in self.doc.sections:
            sec.orientation = WD_ORIENT.LANDSCAPE if s.orientation == "landscape" else WD_ORIENT.PORTRAIT
            sec.page_width, sec.page_height = Inches(w), Inches(h)
            sec.top_margin, sec.bottom_margin = Inches(top), Inches(bottom)
            sec.left_margin, sec.right_margin = Inches(left), Inches(right)
            sec.header_distance = Inches(min(0.5, top / 2))
            sec.footer_distance = Inches(min(0.5, bottom / 2))

    def _scale_images(self) -> None:
        """Shrink inline images that would now be wider than the text area."""
        max_w = min(sec.page_width - sec.left_margin - sec.right_margin for sec in self.doc.sections)
        for inline in self.doc.element.body.xpath(".//wp:inline"):
            if inline.xpath("ancestor::w:tc"):
                continue
            extent = inline.find(qn("wp:extent"))
            if extent is None:
                continue
            cx, cy = int(extent.get("cx", 0)), int(extent.get("cy", 0))
            if cx <= max_w or cx == 0:
                continue
            ratio = max_w / cx
            extent.set("cx", str(int(cx * ratio)))
            extent.set("cy", str(int(cy * ratio)))
            for ext in inline.xpath(".//pic:spPr/a:xfrm/a:ext"):
                ext.set("cx", str(int(int(ext.get("cx", cx)) * ratio)))
                ext.set("cy", str(int(int(ext.get("cy", cy)) * ratio)))
            self.stats["images_scaled"] += 1

    # ---- headers / footers / page numbers -------------------------------------------------------
    def _hf_parts(self, section, kind: str):
        parts = [getattr(section, kind)]
        if section.different_first_page_header_footer:
            parts.append(getattr(section, f"first_page_{kind}"))
        if self.doc.settings.odd_and_even_pages_header_footer:
            parts.append(getattr(section, f"even_page_{kind}"))
        return parts

    def _apply_headers_footers(self) -> None:
        s = self.s
        jobs = {"header": (s.header, "top"), "footer": (s.footer, "bottom")}
        for idx, sec in enumerate(self.doc.sections):
            for kind, (cfg, where) in jobs.items():
                pn_here = s.page_numbers.enabled and s.page_numbers.position.startswith(where)
                if not (cfg.enabled or pn_here):
                    continue
                for part in self._hf_parts(sec, kind):
                    if part.is_linked_to_previous and idx > 0:
                        continue  # inherits from the previous section, which we already handled
                    if part.is_linked_to_previous:
                        part.is_linked_to_previous = False
                    self._fill_part(part, cfg, pn_here)

    def _fill_part(self, part, cfg, add_page_number: bool) -> None:
        s = self.s
        size = min(10, s.font_size)
        family = s.font_family
        if cfg.enabled:
            if any(has_drawing(p._p) for p in part.paragraphs):  # keep logos etc.
                para = part.add_paragraph()
            else:
                for extra in part.paragraphs[1:]:
                    extra._p.getparent().remove(extra._p)
                para = part.paragraphs[0]
                para.clear()
            run = para.add_run(cfg.text)
            _font_run(run, family, size)
            para.alignment = ALIGN[cfg.alignment]
        if add_page_number:
            for p in list(part.paragraphs):  # drop an existing bare page-number paragraph
                if self._has_page_field(p._p) and p.text.strip().isdigit():
                    p._p.getparent().remove(p._p)
            paras = part.paragraphs
            para = paras[-1] if paras and is_blank(paras[-1]._p) and not cfg.enabled else part.add_paragraph()
            para.alignment = ALIGN[s.page_numbers.position.split("-")[1]]
            self._add_page_field(para, family, size)

    @staticmethod
    def _has_page_field(p_el) -> bool:
        return any((i.text or "").strip().upper().startswith("PAGE") for i in p_el.iter(qn("w:instrText"))) or \
               any((f.get(qn("w:instr")) or "").strip().upper().startswith("PAGE") for f in p_el.iter(qn("w:fldSimple")))

    @staticmethod
    def _add_page_field(paragraph, family, size) -> None:
        def run_with(child):
            r = paragraph.add_run()
            _font_run(r, family, size)
            r._r.append(child)
            return r
        begin, sep, end = (OxmlElement("w:fldChar") for _ in range(3))
        begin.set(qn("w:fldCharType"), "begin")
        sep.set(qn("w:fldCharType"), "separate")
        end.set(qn("w:fldCharType"), "end")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = " PAGE "
        run_with(begin)
        run_with(instr)
        run_with(sep)
        paragraph.add_run("1")  # cached result; Word/LibreOffice recalculate it per page
        _font_run(paragraph.runs[-1], family, size)
        run_with(end)

    # ---- reporting ---------------------------------------------------------------------------------
    def _after(self) -> dict:
        s = self.s
        top, bottom, left, right = self._margins()
        return {
            "font": s.font_family if s.normalize_fonts else "Unchanged",
            "size": f"{s.font_size:g} pt" if s.normalize_font_size else "Unchanged",
            "alignment": ALIGN_LABEL[s.alignment] if s.normalize_alignment else "Unchanged",
            "line_spacing": spacing_label(s.line_spacing) if s.normalize_line_spacing else "Unchanged",
            "page": f"{s.page_size} {s.orientation}",
            "margins": f'{top:g}" top, {bottom:g}" bottom, {left:g}" left, {right:g}" right',
        }

    def _summarise(self) -> None:
        s, st, out = self.s, self.stats, self.summary
        if s.normalize_fonts:
            out.append(f"Font normalized to {s.font_family}")
        if s.normalize_font_size:
            out.append(f"Body text set to {s.font_size:g} pt")
        if s.font_style != "normal":
            out.append(f"{s.font_style.capitalize()} applied to body text")
        if s.normalize_alignment:
            out.append(f"Paragraph alignment set to {ALIGN_LABEL[s.alignment]}")
        if s.normalize_line_spacing:
            out.append(f"Line spacing set to {spacing_label(s.line_spacing)}")
        if s.normalize_paragraph_spacing:
            out.append(f"Paragraph spacing set to {s.space_before:g} pt before / {s.space_after:g} pt after")
            if s.first_line_indent:
                out.append(f"First-line indent set to {s.first_line_indent:g} in")
        if st["labels_aligned"]:
            out.append(f"{st['labels_aligned']} chapter/unit labels set to {ALIGN_LABEL[s.chapter_labels.alignment]}")
        if st["centered_kept"]:
            out.append(f"{st['centered_kept']} short centered lines kept centered")
        if st["headings_formatted"]:
            out.append(f"Title and heading styles applied to {st['headings_formatted']} paragraphs")
        else:
            out.append("No Word heading styles found, so heading settings were not used")
        out.append(f"Page size set to {s.page_size} ({s.orientation})")
        out.append(f"Margins applied ({s.margin_preset})")
        if st["blank_removed"]:
            out.append(f"Removed {st['blank_removed']} extra blank paragraphs")
        if st["images_scaled"]:
            out.append(f"Resized {st['images_scaled']} images to fit the text area")
        if s.header.enabled:
            out.append("Header added")
        if s.footer.enabled:
            out.append("Footer added")
        if s.page_numbers.enabled:
            out.append(f"Page numbers added ({s.page_numbers.position.replace('-', ' ')})")
        if st["images_removed"]:
            out.append(f"Removed {st['images_removed']} images (as requested)")
        if st["tables_removed"]:
            out.append(f"Removed {st['tables_removed']} tables (as requested)")
        if st["links_removed"]:
            out.append(f"Converted {st['links_removed']} hyperlinks to plain text (as requested)")
