"""Read-only analysis of a DOCX file. Never modifies the document."""
import io
import math
import re
import zipfile
from collections import Counter

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Length

from models.formatting_settings import PAGE_SIZES_INCHES
from services.docx_helpers import (
    heading_level, is_blank, is_chapter_label, is_special_style, rfonts_name, run_text, runs_of,
    style_chain, theme_fonts,
)

ALIGN_NAMES = {"LEFT": "Left", "CENTER": "Center", "RIGHT": "Right", "JUSTIFY": "Justified",
               "DISTRIBUTE": "Distributed", "BOTH": "Justified"}


def _align_name(value) -> str:
    if value is None:
        return "Left"
    key = str(value).split(".")[-1].split(" ")[0].upper()
    return ALIGN_NAMES.get(key, key.title())


def spacing_label(v) -> str:
    if v is None:
        return "Single"
    if isinstance(v, Length):
        return f"Exactly {v.pt:g} pt"
    v = round(float(v), 2)
    return "Single" if v == 1 else "Double" if v == 2 else f"{v:g}"


class DocumentAnalyzer:
    def __init__(self, data: bytes, file_name: str):
        self.data, self.file_name = data, file_name
        self.doc = Document(io.BytesIO(data))
        self.themes = theme_fonts(self.doc)
        styles_el = self.doc.styles.element
        self._dd_rfonts = styles_el.find(
            "/".join(qn(t) for t in ("w:docDefaults", "w:rPrDefault", "w:rPr", "w:rFonts")))
        self._dd_sz = styles_el.find(
            "/".join(qn(t) for t in ("w:docDefaults", "w:rPrDefault", "w:rPr", "w:sz")))
        self._dd_spacing = styles_el.find(
            "/".join(qn(t) for t in ("w:docDefaults", "w:pPrDefault", "w:pPr", "w:spacing")))

    # ---- effective-value resolution -------------------------------------------------
    def _font(self, run, paragraph) -> str:
        name = rfonts_name(run._r.rPr.rFonts if run._r.rPr is not None else None, self.themes)
        if name:
            return name
        styles = []
        if run._r.rPr is not None and run._r.rPr.rStyle is not None:
            styles += list(style_chain(run.style))
        styles += list(style_chain(paragraph.style))
        for st in styles:
            rpr = st.element.rPr
            name = rfonts_name(rpr.rFonts if rpr is not None else None, self.themes)
            if name:
                return name
        return rfonts_name(self._dd_rfonts, self.themes) or self.themes["minor"] or "Default"

    def _size(self, run, paragraph) -> float:
        if run.font.size:
            return run.font.size.pt
        styles = []
        if run._r.rPr is not None and run._r.rPr.rStyle is not None:
            styles += list(style_chain(run.style))
        styles += list(style_chain(paragraph.style))
        for st in styles:
            if st.font.size:
                return st.font.size.pt
        if self._dd_sz is not None and self._dd_sz.get(qn("w:val")):
            return int(self._dd_sz.get(qn("w:val"))) / 2
        return 10.0

    def _alignment(self, paragraph) -> str:
        if paragraph.alignment is not None:
            return _align_name(paragraph.alignment)
        for st in style_chain(paragraph.style):
            if st.paragraph_format.alignment is not None:
                return _align_name(st.paragraph_format.alignment)
        return "Left"

    def _line_spacing(self, paragraph):
        if paragraph.paragraph_format.line_spacing is not None:
            return paragraph.paragraph_format.line_spacing
        for st in style_chain(paragraph.style):
            if st.paragraph_format.line_spacing is not None:
                return st.paragraph_format.line_spacing
        sp = self._dd_spacing
        if sp is not None and sp.get(qn("w:line")):
            line = int(sp.get(qn("w:line")))
            rule = sp.get(qn("w:lineRule")) or "auto"
            return line / 240 if rule == "auto" else Length(line * 635)
        return 1.0

    # ---- analysis ---------------------------------------------------------------------
    def _pages(self, words: int) -> dict:
        try:
            with zipfile.ZipFile(io.BytesIO(self.data)) as z:
                if "docProps/app.xml" in z.namelist():
                    m = re.search(r"<Pages>(\d+)</Pages>", z.read("docProps/app.xml").decode("utf8", "ignore"))
                    if m and int(m.group(1)) > 0:
                        return {"value": int(m.group(1)), "source": "saved by Word"}
        except Exception:
            pass
        return {"value": max(1, math.ceil(words / 450)), "source": "estimate"}

    def _page_info(self) -> dict:
        s = self.doc.sections[0]
        w, h = (s.page_width or 0) / 914400, (s.page_height or 0) / 914400
        short, long = sorted((w, h))
        name = "Custom"
        for n, (pw, ph) in PAGE_SIZES_INCHES.items():
            if abs(short - pw) < 0.1 and abs(long - ph) < 0.1:
                name = n
        inch = lambda v: round((v or 0) / 914400, 2)
        return {"size": name, "orientation": "landscape" if w > h else "portrait",
                "width_in": round(w, 2), "height_in": round(h, 2),
                "margins": {"top": inch(s.top_margin), "bottom": inch(s.bottom_margin),
                            "left": inch(s.left_margin), "right": inch(s.right_margin)}}

    def analyze(self) -> dict:
        doc, body = self.doc, self.doc.element.body
        paragraphs = doc.paragraphs
        heading_counts, title_found = Counter(), False
        fonts, sizes, aligns, spacings = Counter(), Counter(), Counter(), Counter()
        non_empty = blank = words = chapters = 0
        chapter_example = None

        def count_fonts(par, count_sizes):
            for run in runs_of(par):
                n = len(run_text(run).strip())
                if n:
                    fonts[self._font(run, par)] += n
                    if count_sizes:
                        sizes[self._size(run, par)] += n

        for p in paragraphs:
            if is_blank(p._p):
                blank += 1
                continue
            non_empty += 1
            words += len(p.text.split())
            lvl = heading_level(p)
            if lvl == 0:
                title_found = True
            elif lvl:
                heading_counts[lvl] += 1
            if lvl is not None or is_special_style(p):
                continue
            if is_chapter_label(p.text):
                chapters += 1
                chapter_example = chapter_example or " ".join(p.text.split())
                count_fonts(p, True)
                continue  # labels are centered on purpose, so they don't count as inconsistent
            if p.text.strip():
                aligns[self._alignment(p)] += 1
                spacings[spacing_label(self._line_spacing(p))] += 1
            count_fonts(p, True)

        for table in doc.tables:  # fonts only; table cells keep their own layout
            for p in table._tbl.iter(qn("w:p")):
                from docx.text.paragraph import Paragraph
                count_fonts(Paragraph(p, doc), False)

        images = len(body.xpath(".//pic:pic")) or 0
        images = max(images, len(body.xpath(".//w:drawing")) + len(body.xpath(".//w:pict")))
        total_chars = sum(fonts.values()) or 1
        headings = sum(heading_counts.values())
        overview = {
            "paragraphs": non_empty, "headings": headings, "tables": len(doc.tables),
            "images": images, "hyperlinks": len(body.xpath(".//w:hyperlink")),
            "blank_paragraphs": blank, "words": words, "sections": len(doc.sections),
            "pages": self._pages(words), "has_title": title_found, "chapter_labels": chapters,
            "heading_levels": {str(k): v for k, v in sorted(heading_counts.items())},
        }
        dominant = aligns.most_common(1)[0][0] if aligns else None
        off_align = sum(c for a, c in aligns.items() if a != dominant)
        issues = [{"level": "ok", "text": "DOCX file is valid"}]
        if len(fonts) > 1:
            issues.append({"level": "warn", "text": f"{len(fonts)} different fonts detected"})
        elif fonts:
            issues.append({"level": "ok", "text": f"One font used ({next(iter(fonts))})"})
        if len(sizes) > 1:
            issues.append({"level": "warn", "text": f"{len(sizes)} different body font sizes detected"})
        elif sizes:
            issues.append({"level": "ok", "text": f"One body font size used ({next(iter(sizes)):g} pt)"})
        if off_align:
            issues.append({"level": "warn", "text":
                           f"{off_align} body paragraphs use a different alignment than the most common one ({dominant})"})
        elif aligns:
            issues.append({"level": "ok", "text": f"Body alignment is consistent ({dominant})"})
        if len(spacings) > 1:
            issues.append({"level": "warn", "text": f"{len(spacings)} different line spacings detected"})
        if blank:
            issues.append({"level": "warn", "text": f"{blank} blank paragraphs detected"})
        if headings == 0:
            issues.append({"level": "warn", "text": "No Word heading styles found, so heading settings will have no effect"})
        if chapters:
            issues.append({"level": "ok", "text": f"{chapters} chapter/unit label{'s' if chapters != 1 else ''} detected (for example \u201c{chapter_example}\u201d)"})
        if overview["tables"]:
            issues.append({"level": "ok", "text": f"{overview['tables']} table{'s' if overview['tables'] != 1 else ''} detected"})
        if images:
            issues.append({"level": "ok", "text": f"{images} image{'s' if images != 1 else ''} detected"})
        if non_empty == 0 and not images and not overview["tables"]:
            issues.append({"level": "warn", "text": "The document appears to be empty"})

        return {
            "file_name": self.file_name, "file_size": len(self.data), "overview": overview,
            "fonts": [{"name": n, "chars": c, "percent": round(100 * c / total_chars)} for n, c in fonts.most_common()],
            "font_sizes": [{"size": s, "chars": c} for s, c in sorted(sizes.items(), key=lambda x: -x[1])],
            "alignments": [{"name": n, "count": c} for n, c in aligns.most_common()],
            "line_spacing": [{"label": n, "count": c} for n, c in spacings.most_common()],
            "page": self._page_info(), "issues": issues,
            "before": {
                "fonts": len(fonts), "font_sizes": len(sizes), "alignments": len(aligns),
                "line_spacing": "Inconsistent" if len(spacings) > 1 else (next(iter(spacings), "—")),
            },
        }
