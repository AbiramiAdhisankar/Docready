"""Generate a messy sample DOCX for development/testing.

Usage: python scripts/make_sample.py [output_path]
The file contains a title, headings, mixed fonts/sizes/alignments, a table, an image,
a hyperlink and blank paragraphs. It is test data only, never user data.
"""
import io
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image


def add_hyperlink(paragraph, url, text):
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def build(path: str) -> None:
    doc = Document()
    doc.add_paragraph("Smart Campus Seminar Report", style="Title")
    doc.add_paragraph()
    for label in ("UNIT I", "INTRODUCTION"):  # chapter-style lines that must stay centered
        doc.add_paragraph(label).alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading("1. Introduction", level=1)
    lorem = ("Reinforcement learning studies how an agent improves its decisions through interaction "
             "with an environment. This sentence is repeated to give the paragraph some body. ") * 3
    p = doc.add_paragraph()
    r = p.add_run(lorem); r.font.name = "Comic Sans MS"; r.font.size = Pt(14)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.add_paragraph()
    doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run(lorem); r.font.name = "Courier New"; r.font.size = Pt(10)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(lorem)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p = doc.add_paragraph("Read more at ")
    add_hyperlink(p, "https://example.com/docs", "the project page")
    doc.add_heading("2. Method", level=1)
    doc.add_heading("2.1 Value iteration", level=2)
    p = doc.add_paragraph(lorem)
    p.runs[0].font.name = "Verdana"; p.runs[0].font.size = Pt(11)

    table = doc.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    for i, (a, b) in enumerate([("Algorithm", "Update rule"), ("Value iteration", "Bellman optimality"),
                                ("TD(0)", "Bootstrapped estimate")]):
        table.cell(i, 0).text, table.cell(i, 1).text = a, b
    doc.add_paragraph()  # blank paragraph after table (must survive)

    img = io.BytesIO()
    Image.new("RGB", (900, 450), (60, 110, 200)).save(img, "PNG")
    img.seek(0)
    doc.add_picture(img, width=Inches(6))
    doc.add_paragraph()
    doc.add_paragraph("Figure 1 caption text.", style="Caption")
    doc.add_heading("3. Conclusion", level=1)
    p = doc.add_paragraph("Short closing paragraph in the default font.")
    doc.save(path)


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else "sample.docx")
